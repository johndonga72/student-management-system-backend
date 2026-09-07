from typing import Any
from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook


class AttendanceExcelReader:
    """
    Reads attendance data from an Excel workbook.

    This class is responsible only for reading Excel data.
    Validation and database operations are handled elsewhere.
    """

    SHEET_NAME = "Attendance"

    REQUIRED_COLUMNS = (
        "student_number",
        "employee_id",
        "subject_code",
        "attendance_date",
        "status",
    )

    OPTIONAL_COLUMNS = (
        "remarks",
    )

    @classmethod
    def read(cls, uploaded_file) -> list[dict[str, Any]]:
        """
        Read attendance rows from an uploaded Excel file.

        Args:
            uploaded_file: Django UploadedFile object.

        Returns:
            A list of dictionaries representing Excel rows.

        Raises:
            ValueError: If the expected worksheet does not exist.
        """

        workbook = cls._load_workbook(uploaded_file)

        try:
            worksheet = workbook[cls.SHEET_NAME]
        except KeyError as exc:
            raise ValueError(
                f"Excel sheet '{cls.SHEET_NAME}' does not exist."
            ) from exc

        return cls._read_rows(worksheet)

    @staticmethod
    def _load_workbook(uploaded_file) -> Workbook:
        """
        Load the uploaded Excel workbook.

        data_only=True returns stored cell values instead of formulas.
        """

        return load_workbook(
            filename=uploaded_file,
            read_only=True,
            data_only=True,
        )

    @classmethod
    def _read_rows(cls, worksheet) -> list[dict[str, Any]]:
        """
        Convert worksheet rows into dictionaries.
        """

        rows = worksheet.iter_rows(
            values_only=True,
        )

        headers = next(rows, None)

        if not headers:
            return []

        normalized_headers = [
            cls._normalize_header(header)
            for header in headers
        ]

        attendance_rows = []

        for row in rows:
            if cls._is_empty_row(row):
                continue

            row_data = dict(
                zip(normalized_headers, row)
            )

            attendance_rows.append(row_data)

        return attendance_rows

    @staticmethod
    def _normalize_header(header: Any) -> str:
        """
        Normalize an Excel column header.
        """

        if header is None:
            return ""

        return str(header).strip().lower()

    @staticmethod
    def _is_empty_row(row: tuple[Any, ...]) -> bool:
        """
        Check whether an Excel row contains no data.
        """

        return all(
            value is None
            for value in row
        )